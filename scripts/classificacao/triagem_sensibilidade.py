"""Triagem automática de sensibilidade — Monitora e PANs (apoio à Fase 3 do roadmap).

Este script NÃO decide a sensibilidade de nenhum documento — essa decisão é
humana e da equipe do RAN (ver docs/processos/classificacao_sensibilidade.md).
Ele varre o texto já extraído (07_processados/textos_extraidos/) procurando
padrões que merecem atenção — coordenadas, dados pessoais, palavras de
restrição — e gera uma planilha priorizada, para que a revisão manual comece
pelos documentos com maior chance de conter algo sensível, em vez de percorrer
os 815 PDFs em ordem alfabética.

Cada linha da planilha de saída tem colunas em branco (decisao, revisado_por,
data_revisao, observacoes) para a equipe do RAN preencher durante a revisão.

Uso:
    python scripts/classificacao/triagem_sensibilidade.py
    python scripts/classificacao/triagem_sensibilidade.py --fonte monitora

Saída:
    06_inventario/triagem_sensibilidade_monitora_pans.xlsx
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[2]
TEXTOS_DIR = RAIZ / "07_processados" / "textos_extraidos"
SAIDA = RAIZ / "06_inventario" / "triagem_sensibilidade_monitora_pans.xlsx"

FONTES = ["monitora", "pans"]

sys.path.insert(0, str(Path(__file__).resolve().parent))
from regras_pii import FLAGS  # noqa: E402  (padrões compartilhados com gerar_chunks.py)

ORDEM_PRIORIDADE = {"ALTA": 0, "MEDIA": 1, "BAIXA": 2}
COR_PRIORIDADE = {
    "ALTA": "F8CBAD",
    "MEDIA": "FFE699",
    "BAIXA": "C6E0B4",
}


def contexto(texto: str, inicio: int, fim: int, largura: int = 60) -> str:
    a = max(0, inicio - largura)
    b = min(len(texto), fim + largura)
    trecho = texto[a:b].replace("\n", " ").strip()
    return f"(...) {trecho} (...)"


def analisar_documento(doc: dict) -> list[dict]:
    ocorrencias = []
    paginas = doc.get("paginas") or [{"pagina": None, "texto": doc.get("texto_completo", "")}]
    for pag in paginas:
        texto = pag.get("texto") or ""
        num_pagina = pag.get("pagina")
        for chave, regex, severidade, descricao, _redigivel, _placeholder in FLAGS:
            for m in regex.finditer(texto):
                ocorrencias.append(
                    {
                        "flag": chave,
                        "severidade": severidade,
                        "descricao": descricao,
                        "pagina": num_pagina,
                        "trecho": contexto(texto, m.start(), m.end()),
                    }
                )
    return ocorrencias


def resumir_documento(doc: dict, ocorrencias: list[dict]) -> dict:
    if ocorrencias:
        prioridade = min((o["severidade"] for o in ocorrencias), key=lambda s: ORDEM_PRIORIDADE[s])
    else:
        prioridade = "BAIXA"

    flags_unicos = {}
    for o in ocorrencias:
        flags_unicos.setdefault(o["flag"], {"descricao": o["descricao"], "contagem": 0})
        flags_unicos[o["flag"]]["contagem"] += 1
    flags_resumo = "; ".join(
        f"{info['descricao']} ({info['contagem']}x)" for info in flags_unicos.values()
    )

    ocorrencias_ordenadas = sorted(ocorrencias, key=lambda o: ORDEM_PRIORIDADE[o["severidade"]])
    exemplos = []
    for o in ocorrencias_ordenadas[:5]:
        pag = f"p.{o['pagina']}" if o["pagina"] is not None else "p.?"
        exemplos.append(f"[{pag} | {o['descricao']}] {o['trecho']}")
    trechos_resumo = "\n".join(exemplos)

    return {
        "fonte": doc.get("fonte"),
        "subpasta": doc.get("subdir") or "",
        "arquivo": doc.get("nome_arquivo"),
        "url_origem": doc.get("url_origem") or "",
        "num_paginas": doc.get("num_paginas"),
        "provavelmente_escaneado": doc.get("provavelmente_escaneado"),
        "prioridade": prioridade,
        "total_ocorrencias": len(ocorrencias),
        "flags_encontrados": flags_resumo,
        "trechos_exemplo": trechos_resumo,
        "decisao": "",
        "revisado_por": "",
        "data_revisao": "",
        "observacoes": "",
    }


def coletar_documentos(fontes: list[str]):
    for fonte in fontes:
        pasta = TEXTOS_DIR / fonte
        if not pasta.exists():
            print(f"  [{fonte}] {pasta} não encontrada — pulando.", file=sys.stderr)
            continue
        for caminho in sorted(pasta.rglob("*.json")):
            import json

            try:
                doc = json.loads(caminho.read_text(encoding="utf-8"))
            except Exception as e:
                print(f"  ✗ {caminho} — ERRO ao ler: {e}", file=sys.stderr)
                continue
            yield doc


COLUNAS = [
    ("fonte", 12),
    ("subpasta", 22),
    ("arquivo", 40),
    ("url_origem", 40),
    ("num_paginas", 10),
    ("provavelmente_escaneado", 10),
    ("prioridade", 10),
    ("total_ocorrencias", 10),
    ("flags_encontrados", 45),
    ("trechos_exemplo", 70),
    ("decisao", 16),
    ("revisado_por", 16),
    ("data_revisao", 14),
    ("observacoes", 30),
]


def sanitizar(valor):
    """Remove caracteres de controle que o Excel/openpyxl rejeita (comuns em PDFs)."""
    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)
    return valor


def gerar_planilha(linhas: list[dict], saida: Path):
    linhas_ordenadas = sorted(
        linhas, key=lambda r: (ORDEM_PRIORIDADE[r["prioridade"]], r["fonte"], r["subpasta"], r["arquivo"])
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Triagem"

    for col_idx, (nome, largura) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col_idx, value=nome)
        cel.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = largura
    ws.freeze_panes = "A2"

    for row_idx, linha in enumerate(linhas_ordenadas, start=2):
        for col_idx, (nome, _largura) in enumerate(COLUNAS, start=1):
            cel = ws.cell(row=row_idx, column=col_idx, value=sanitizar(linha.get(nome)))
            cel.alignment = Alignment(vertical="top", wrap_text=nome in ("trechos_exemplo", "flags_encontrados"))
        cor = COR_PRIORIDADE.get(linha["prioridade"])
        if cor:
            ws.cell(row=row_idx, column=COLUNAS.index(("prioridade", 10)) + 1).fill = PatternFill(
                "solid", fgColor=cor
            )

    # Aba de resumo
    ws2 = wb.create_sheet("Resumo")
    contagem = {"ALTA": 0, "MEDIA": 0, "BAIXA": 0}
    for linha in linhas_ordenadas:
        contagem[linha["prioridade"]] += 1
    ws2.append(["Prioridade", "Documentos", "Observação"])
    ws2.cell(row=1, column=1).font = Font(bold=True)
    ws2.cell(row=1, column=2).font = Font(bold=True)
    ws2.cell(row=1, column=3).font = Font(bold=True)
    ws2.append(["ALTA", contagem["ALTA"], "Revisar primeiro — CPF/e-mail/coordenada/palavra de restrição"])
    ws2.append(["MEDIA", contagem["MEDIA"], "Revisar em seguida — menção genérica a localização/dado pessoal"])
    ws2.append(["BAIXA", contagem["BAIXA"], "Nenhum padrão encontrado — recomenda-se amostragem, não pular"])
    ws2.append(["TOTAL", len(linhas_ordenadas), ""])
    for col, largura in [("A", 14), ("B", 12), ("C", 65)]:
        ws2.column_dimensions[col].width = largura

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--fonte", choices=FONTES, default=None, help="Rodar só para esta fonte (padrão: todas)")
    args = parser.parse_args()

    fontes = [args.fonte] if args.fonte else FONTES

    linhas = []
    for doc in coletar_documentos(fontes):
        ocorrencias = analisar_documento(doc)
        linhas.append(resumir_documento(doc, ocorrencias))

    if not linhas:
        print("Nenhum documento encontrado em 07_processados/textos_extraidos/ — rode a extração antes.")
        return

    gerar_planilha(linhas, SAIDA)

    contagem = {"ALTA": 0, "MEDIA": 0, "BAIXA": 0}
    for linha in linhas:
        contagem[linha["prioridade"]] += 1

    print(f"Triagem concluída — {len(linhas)} documentos analisados.")
    print(f"  ALTA prioridade:  {contagem['ALTA']}")
    print(f"  MÉDIA prioridade: {contagem['MEDIA']}")
    print(f"  BAIXA prioridade: {contagem['BAIXA']}")
    print(f"\nPlanilha salva em {SAIDA.relative_to(RAIZ)}")
    print("\nLembrete: este script só sinaliza padrões — a classificação final")
    print("(autorizado/sensível) é decisão humana da equipe do RAN, coluna 'decisao'.")


if __name__ == "__main__":
    main()
